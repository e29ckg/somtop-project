from pathlib import Path
from zipfile import ZipFile
from docx import Document

REFERENCE = Path(r"C:\Users\PKKJC-IT-519\Downloads\ใบเปลี่ยนเวรพ.สมทบ.docx")
OUTPUT = Path(r"C:\project\somtop-project\backend\templates\duty_swap_template.docx")


def replace_paragraph(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


doc = Document(REFERENCE)
paragraphs = doc.paragraphs

# The supplied file contains floating scan/signature layers positioned relative
# to paragraphs. They are not visible in the clean reference page, but moving
# the anchor paragraphs makes parts of those layers appear below the page.
# Remove the hidden drawing layers while retaining the document's native text,
# spacing, borders, checkboxes, and page setup.
for drawing in doc.element.xpath(".//w:drawing"):
    drawing.getparent().remove(drawing)

replace_paragraph(paragraphs[2], "ที่\t-\t\t\tวันที่ {request_date_th}")
replace_paragraph(
    paragraphs[6],
    "\tตามคำสั่งศาลเยาวชนและครอบครัวจังหวัดประจวบคีรีขันธ์ ที่ {order_number} "
    "ประจำเดือน {order_month_th} เรื่อง ให้ผู้พิพากษาสมทบมาปฏิบัติหน้าที่เป็นองค์คณะและรับรายงานตัว "
    "โดยให้ข้าพเจ้า {requester_name} ผู้พิพากษาสมทบ ปฏิบัติหน้าที่ดังกล่าวในวันที่ {duty_date_th} นั้น",
)
replace_paragraph(
    paragraphs[8],
    "\t\tเนื่องจากข้าพเจ้า {requester_name} ผู้พิพากษาสมทบ มีเหตุจำเป็น {reason} "
    "ไม่สามารถมาปฏิบัติหน้าที่ในวันดังกล่าวได้ จึงมีความประสงค์ขอเปลี่ยนเวรปฏิบัติหน้าที่ดังกล่าว "
    "โดยให้ {replacement_name} ผู้พิพากษาสมทบ มาปฏิบัติหน้าที่ในวันดังกล่าวแทนข้าพเจ้า "
    "ซึ่ง {replacement_name} ตกลงยินยอมด้วยแล้ว",
)
replace_paragraph(paragraphs[13], "      ลงชื่อ........................................................ผู้ขอเปลี่ยนเวร")
replace_paragraph(paragraphs[14], "                  ({requester_name})")
replace_paragraph(paragraphs[16], "\t\t\t\t\t    ลงชื่อ.......................................................ผู้รับเปลี่ยนเวร")
replace_paragraph(paragraphs[17], "\t\t\t\t\t\t           ({replacement_name})")
replace_paragraph(paragraphs[23], "({approver_name})")
replace_paragraph(paragraphs[24], "{approver_title}")

doc.save(OUTPUT)

with ZipFile(REFERENCE) as before, ZipFile(OUTPUT) as after:
    preserved = {
        name for name in before.namelist()
        if name.startswith("customXml/")
    }
    missing = preserved.difference(after.namelist())
    if missing:
        raise RuntimeError(f"Preserve-only parts missing: {sorted(missing)}")
