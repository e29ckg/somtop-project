from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

REFERENCE = Path(r"C:\Users\PKKJC-IT-519\Downloads\คำสั่งการอยู่เวรผู้พิพากษาสมทบ.docx")
OUTPUT = Path(r"C:\project\somtop-project\backend\templates\duty_order_template.docx")


def replace_paragraph(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


doc = Document(REFERENCE)
paragraphs = doc.paragraphs

replace_paragraph(paragraphs[4], "คำสั่ง{court_name}")
replace_paragraph(paragraphs[5], "ที่ {order_number}")
paragraphs[5].alignment = WD_ALIGN_PARAGRAPH.CENTER
replace_paragraph(paragraphs[6], "เรื่อง {order_title}")
replace_paragraph(paragraphs[7], "ประจำเดือน {order_month_th} เวลา {start_time} - {end_time} นาฬิกา")

# Use the source schedule paragraph as the repeating record so its tabs,
# indents, spacing, and font remain the visual authority.
if paragraphs[10]._p.pPr is not None:
    paragraphs[10]._p.remove(paragraphs[10]._p.pPr)
if paragraphs[9]._p.pPr is not None:
    paragraphs[10]._p.insert(0, deepcopy(paragraphs[9]._p.pPr))
replace_paragraph(paragraphs[9], "{#schedules}")
replace_paragraph(paragraphs[10], "{duty_date_th}\t{team_name}\t{names}")
replace_paragraph(paragraphs[11], "{/schedules}")

replace_paragraph(paragraphs[51], "หมายเหตุ * หมายถึง ผู้ได้รับการอบรมและแต่งตั้งเป็นผู้ประสานการประชุมของศาล")
replace_paragraph(paragraphs[53], "อนึ่ง หากท่านไม่สามารถมาปฏิบัติหน้าที่ตามคำสั่งได้ ให้จัดหาผู้ปฏิบัติหน้าที่แทนเป็นลายลักษณ์อักษร")
replace_paragraph(paragraphs[54], "สั่ง ณ วันที่ {issued_date_th}")
replace_paragraph(paragraphs[56], "({chief_judge_name})")
replace_paragraph(paragraphs[57], "{chief_judge_position}")
replace_paragraph(paragraphs[58], "{chief_judge_acting_position}")

for index in range(50, 11, -1):
    remove_paragraph(paragraphs[index])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
